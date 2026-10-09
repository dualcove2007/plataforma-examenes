from io import BytesIO

from django.conf import settings
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from apps.attempts.models import ResultadoExamen

R = ResultadoExamen.Revision
_INICIOS_PELIGROSOS = ("=", "+", "-", "@", "\t", "\r")
_ENCABEZADOS = [
    "Estudiante",
    "Correo",
    "Intento",
    "Estado del intento",
    "Puntaje",
    "Puntaje máximo",
    "Nota (0-5)",
    "Revisión",
    "Resultado",
]
_ANCHOS = [30, 32, 9, 18, 10, 15, 11, 14, 22]


def _escribir(ws, fila, valores, negrita=False, relleno=None):
    for col, valor in enumerate(valores, start=1):
        celda = ws.cell(row=fila, column=col, value=valor)
        # Un texto que empiece por "=" se guarda como texto, nunca como fórmula
        if isinstance(valor, str) and valor[:1] in _INICIOS_PELIGROSOS:
            celda.data_type = "s"
        if negrita:
            celda.font = Font(bold=True)
        if relleno:
            celda.fill = relleno


def _veredicto(resultado):
    if resultado.estado_revision != R.REVISADO:
        return "Pendiente de revisión"
    aprobado = resultado.nota_final >= settings.NOTA_APROBATORIA
    return "Aprobado" if aprobado else "Reprobado"


def generar_excel_notas(examen):
    """Devuelve los bytes de un .xlsx con las notas de todos los intentos del examen."""
    resultados = list(
        ResultadoExamen.objects.filter(intento__asignacion__examen=examen)
        .select_related("intento__asignacion__estudiante")
        .order_by("intento__asignacion__estudiante__nombre", "intento__numero_intento")
    )
    revisados = [r.nota_final for r in resultados if r.estado_revision == R.REVISADO]
    promedio = round(sum(revisados) / len(revisados), 1) if revisados else "—"

    wb = Workbook()
    ws = wb.active
    ws.title = "Notas"

    _escribir(ws, 1, ["Examen", examen.titulo])
    _escribir(ws, 2, ["Materia", examen.materia.nombre])
    _escribir(ws, 3, ["Nota mínima para aprobar", settings.NOTA_APROBATORIA])
    _escribir(ws, 4, ["Promedio (solo revisados)", promedio])
    for fila in range(1, 5):
        ws.cell(row=fila, column=1).font = Font(bold=True)

    azul = PatternFill("solid", start_color="D9E2F3")
    _escribir(ws, 6, _ENCABEZADOS, negrita=True, relleno=azul)

    for i, r in enumerate(resultados, start=7):
        intento = r.intento
        estudiante = intento.asignacion.estudiante
        _escribir(
            ws,
            i,
            [
                estudiante.nombre,
                estudiante.email,
                intento.numero_intento,
                intento.get_estado_display(),
                r.puntaje_total,
                r.puntaje_maximo,
                r.nota_final,
                r.get_estado_revision_display(),
                _veredicto(r),
            ],
        )
        ws.cell(row=i, column=7).number_format = "0.0"
        ws.cell(row=i, column=7).alignment = Alignment(horizontal="center")

    for col, ancho in enumerate(_ANCHOS, start=1):
        ws.column_dimensions[chr(64 + col)].width = ancho
    ws.freeze_panes = "A7"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()