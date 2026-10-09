from datetime import timedelta
from io import BytesIO

from django.core.management import call_command
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.attempts.models import Intento, ResultadoExamen
from apps.audit.models import Auditoria
from apps.authentication.models import Usuario
from apps.exams.models import AsignacionExamen, Examen
from apps.question_banks.models import Materia

URL = "/api/reportes/notas/"


class ReporteNotasTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("seed", verbosity=0)
        cls.admin = Usuario.objects.get(email="admin@demo.com")
        cls.docente = Usuario.objects.get(email="docente@demo.com")
        cls.estudiante = Usuario.objects.get(email="estudiante@demo.com")
        cls.otro_docente = Usuario.objects.create_user(
            "otro@demo.com", "Otro Docente", "Demo12345*", rol=cls.docente.rol
        )
        materia = Materia.objects.create(nombre="Redes")
        ahora = timezone.now()
        cls.examen = Examen.objects.create(
            docente=cls.docente, materia=materia, titulo="Parcial 1",
            fecha_inicio=ahora - timedelta(days=2), fecha_fin=ahora - timedelta(days=1),
            duracion_minutos=30, estado=Examen.Estado.CERRADO,
        )
        asignacion = AsignacionExamen.objects.create(
            examen=cls.examen, estudiante=cls.estudiante
        )
        intento = Intento.objects.create(
            asignacion=asignacion, numero_intento=1,
            fecha_limite=ahora - timedelta(days=1), estado=Intento.Estado.FINALIZADO,
        )
        ResultadoExamen.objects.create(
            intento=intento, puntaje_total=4, puntaje_maximo=5, nota_final=4.0,
            estado_revision=ResultadoExamen.Revision.REVISADO,
        )

    def get(self, usuario, **params):
        cliente = APIClient()
        if usuario:
            cliente.force_authenticate(usuario)
        return cliente.get(URL, params)

    def test_docente_descarga_excel_de_su_examen(self):
        r = self.get(self.docente, examen=self.examen.pk)
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        self.assertIn(f"notas_examen_{self.examen.pk}.xlsx", r["Content-Disposition"])
        hoja = load_workbook(BytesIO(r.content)).active
        filas = list(hoja.iter_rows(values_only=True))
        fila = next(f for f in filas if f[1] == "estudiante@demo.com")
        self.assertEqual(fila[6], 4.0)
        self.assertEqual(fila[8], "Aprobado")

    def test_admin_puede_descargar_cualquier_examen(self):
        self.assertEqual(self.get(self.admin, examen=self.examen.pk).status_code, 200)

    def test_docente_ajeno_no_ve_el_examen(self):
        self.assertEqual(self.get(self.otro_docente, examen=self.examen.pk).status_code, 404)

    def test_estudiante_no_tiene_permiso(self):
        self.assertEqual(self.get(self.estudiante, examen=self.examen.pk).status_code, 403)

    def test_sin_autenticar(self):
        self.assertIn(self.get(None, examen=self.examen.pk).status_code, (401, 403))

    def test_examen_obligatorio_y_valido(self):
        self.assertEqual(self.get(self.docente).status_code, 400)
        self.assertEqual(self.get(self.docente, examen="abc").status_code, 400)

    def test_queda_registrado_en_auditoria(self):
        self.get(self.docente, examen=self.examen.pk)
        self.assertTrue(
            Auditoria.objects.filter(
                accion="exportar", entidad="Examen", entidad_id=str(self.examen.pk)
            ).exists()
        )