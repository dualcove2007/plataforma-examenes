import io
import re

from django.db import transaction
from django.http import HttpResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from rest_framework.exceptions import ValidationError
from apps.audit.services import registrar

from .serializers import PreguntaSerializer

MAX_OPCIONES = 6
MAX_FILAS = 500
COLUMNAS = (
    ["enunciado", "tipo", "dificultad", "puntaje"]
    + [f"opcion_{i}" for i in range(1, MAX_OPCIONES + 1)]
    + ["correctas"]
)
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def respuesta_excel(wb, nombre):
    buffer = io.BytesIO()
    wb.save(buffer)
    respuesta = HttpResponse(buffer.getvalue(), content_type=XLSX)
    respuesta["Content-Disposition"] = f'attachment; filename="{nombre}"'
    return respuesta


def _libro_con_encabezados():
    wb = Workbook()
    ws = wb.active
    ws.title = "Preguntas"
    ws.append(COLUMNAS)
    for celda in ws[1]:
        celda.font = Font(bold=True)
    return wb, ws


def generar_plantilla():
    wb, ws = _libro_con_encabezados()
    vacias = [""] * MAX_OPCIONES
    ws.append(
        ["Que comando consulta datos?", "opcion_unica", "facil", 1,
         "SELECT", "DROP", *vacias[2:], "1"]
    )
    ws.append(
        ["SQL es un lenguaje declarativo", "verdadero_falso", "media", 1,
         "Verdadero", "Falso", *vacias[2:], "1"]
    )
    ws.append(
        ["Explique que es una llave foranea", "abierta", "dificil", 3,
         *vacias, ""]
    )
    return wb


def exportar_preguntas(queryset):
    wb, ws = _libro_con_encabezados()
    for p in queryset:
        opciones = sorted(p.opciones.all(), key=lambda o: o.id)[:MAX_OPCIONES]
        textos = [o.texto for o in opciones]
        textos += [""] * (MAX_OPCIONES - len(textos))
        correctas = ",".join(str(i + 1) for i, o in enumerate(opciones) if o.es_correcta)
        ws.append([p.enunciado, p.tipo, p.dificultad, p.puntaje, *textos, correctas])
    return wb


def _celda(fila, indices, nombre):
    posicion = indices[nombre]
    valor = fila[posicion] if posicion < len(fila) else None
    return "" if valor is None else str(valor).strip()


def importar_preguntas(archivo, banco, request):
    """Valida TODAS las filas; si alguna falla no se guarda nada."""
    try:
        wb = load_workbook(archivo, read_only=True, data_only=True)
    except Exception:  # zip corrupto, formato no soportado, etc.
        raise ValidationError({"archivo": "No es un Excel (.xlsx) válido."})

    filas = list(wb.active.iter_rows(values_only=True))
    if not filas:
        raise ValidationError({"archivo": "El archivo está vacío."})

    encabezados = [str(c).strip().lower() if c is not None else "" for c in filas[0]]
    faltantes = [c for c in COLUMNAS if c not in encabezados]
    if faltantes:
        raise ValidationError(
            {"archivo": f"Faltan columnas: {', '.join(faltantes)}. Descarga la plantilla."}
        )
    indices = {nombre: encabezados.index(nombre) for nombre in COLUMNAS}

    datos = filas[1:]
    if len(datos) > MAX_FILAS:
        raise ValidationError({"archivo": f"Máximo {MAX_FILAS} filas por archivo."})

    errores, validados = {}, []
    for numero, fila in enumerate(datos, start=2):
        if all(c is None or str(c).strip() == "" for c in fila):
            continue
        correctas = {
            int(n) for n in re.findall(r"\d+", _celda(fila, indices, "correctas"))
        }
        opciones = []
        for i in range(1, MAX_OPCIONES + 1):
            texto = _celda(fila, indices, f"opcion_{i}")
            if texto:
                opciones.append({"texto": texto, "es_correcta": i in correctas})

        serializer = PreguntaSerializer(
            data={
                "banco": banco.id,
                "enunciado": _celda(fila, indices, "enunciado"),
                "tipo": _celda(fila, indices, "tipo").lower(),
                "dificultad": _celda(fila, indices, "dificultad").lower() or "media",
                "puntaje": (_celda(fila, indices, "puntaje") or "1").replace(",", "."),
                "opciones": opciones,
            },
            context={"request": request},
        )
        if serializer.is_valid():
            validados.append(serializer)
        else:
            errores[f"fila {numero}"] = serializer.errors

    if errores:
        raise ValidationError(errores)
    if not validados:
        raise ValidationError({"archivo": "No hay preguntas para importar."})

    with transaction.atomic():
        for serializer in validados:
            serializer.save()
        registrar(
            "importar", "BancoPreguntas", banco.pk, {"preguntas": len(validados)}
        )
    return len(validados)